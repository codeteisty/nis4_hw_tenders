from __future__ import annotations

import io
import json
import os
import re
import ssl
import time
import zipfile
from http.cookiejar import CookieJar
from pathlib import Path
from urllib.parse import urljoin
from urllib.request import HTTPSHandler, HTTPCookieProcessor, Request, build_opener

import openpyxl
from docx import Document
from lxml import html as lxml_html
from pypdf import PdfReader


BASE = "https://zakupki.gov.ru"
ROOT = Path(os.environ.get("WORK_DIR", "work/research"))
CERT = Path(os.environ.get("EIS_CERT_FILE", "certs/russian-trusted-root-ca.pem"))
INVENTORY = ROOT / "document_inventory.json"

SELECTED = [
    "32616181297",
    "32616201339",
    "32616237698",
    "32616257855",
    "32616266570",
    "32616322212",
    "32616269487",
    "32616224459",
    "32616238969",
    "32616240823",
    "0848300048226000330",
    "0335200014926003234",
    "0335200014926003280",
    "0335200014926003283",
    "0372100021226000003",
]


def safe_name(value: str) -> str:
    value = value.replace("\u00a0", " ").strip()
    value = re.sub(r"[\\/:*?\"<>|]+", "_", value)
    return value[:180] or "file"


def request_bytes(opener, url: str) -> bytes:
    request = Request(url, headers={
        "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 Chrome/139.0 Safari/537.36"
    })
    with opener.open(request, timeout=60) as response:
        content = response.read()
    time.sleep(0.75)
    return content


def fetch_html(opener, url: str, out: Path) -> str:
    if out.exists() and out.stat().st_size > 1000:
        content = out.read_bytes()
    else:
        content = request_bytes(opener, url)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_bytes(content)
    return content.decode("utf-8", errors="replace")


def html_text(value: str) -> str:
    tree = lxml_html.fromstring(value)
    for node in tree.xpath("//script|//style|//svg"):
        node.drop_tree()
    return "\n".join(line.strip() for line in tree.text_content().splitlines() if line.strip())


def attachment_rows(value: str) -> list[dict[str, str]]:
    tree = lxml_html.fromstring(value)
    rows: list[dict[str, str]] = []
    seen: set[str] = set()
    for anchor in tree.xpath('//a[contains(@href, "/filestore/public/1.0/download/")]'):
        href = urljoin(BASE, anchor.get("href", ""))
        if not href or href in seen:
            continue
        seen.add(href)
        title = anchor.get("title") or anchor.get("data-tooltip") or " ".join(anchor.text_content().split())
        if "custom-tooltiptext" in title:
            title = " ".join(lxml_html.fromstring(title).text_content().split())
        rows.append({"name": safe_name(title), "url": href})
    return rows


def select_223_files(attachments: list[dict[str, str]]) -> list[dict[str, str]]:
    usable = [
        item for item in attachments
        if not re.search(r"\.(?:part\d+\.rar|rar\.z\d+|\d{3})$", item["name"], flags=re.I)
    ]
    chosen: list[dict[str, str]] = []
    seen_urls: set[str] = set()

    groups = [
        r"проект.{0,20}договор|договор.{0,20}проект|умный договор|том 4",
        r"документац|части? i+[-\u2013]iii|том 1|(^|\W)кд[_ ]|извещен",
        r"техничес|(^|\W)тз(\W|$)|описание.{0,20}(предмет|объект)",
        r"условия выбора финансов|независим.{0,10}гарант|банковск.{0,10}гарант|(^|\W)бг(\W|$)|обеспеч",
        r"обоснован|(^|\W)нмц",
    ]
    for pattern in groups:
        for item in reversed(usable):
            if item["url"] not in seen_urls and re.search(pattern, item["name"], flags=re.I):
                chosen.append(item)
                seen_urls.add(item["url"])
                break

    protocols: list[dict[str, str]] = []
    protocol_names: set[str] = set()
    for item in reversed(usable):
        if not re.search(r"протокол|итог", item["name"], flags=re.I):
            continue
        normalized = re.sub(r"\s+", " ", item["name"].lower())
        if normalized in protocol_names or item["url"] in seen_urls:
            continue
        protocols.append(item)
        protocol_names.add(normalized)
        seen_urls.add(item["url"])
        if len(protocols) == 3:
            break
    chosen.extend(protocols)
    return chosen[:8]


def extract_docx(path: Path) -> str:
    doc = Document(path)
    blocks: list[str] = []
    for paragraph in doc.paragraphs:
        value = paragraph.text.strip()
        if value:
            blocks.append(value)
    for table in doc.tables:
        for row in table.rows:
            values = [" ".join(cell.text.split()) for cell in row.cells]
            if any(values):
                blocks.append(" | ".join(values))
    return "\n".join(blocks)


def extract_pdf(path: Path) -> str:
    reader = PdfReader(path)
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def extract_xlsx(path: Path) -> str:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    rows: list[str] = []
    for sheet in workbook.worksheets:
        rows.append(f"ЛИСТ {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                rows.append(" | ".join(values))
    return "\n".join(rows)


def extract_zip(path: Path) -> str:
    parts: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for member in archive.infolist():
            name = safe_name(member.filename)
            suffix = Path(name).suffix.lower()
            if member.is_dir() or member.file_size > 30_000_000:
                continue
            data = archive.read(member)
            try:
                if suffix == ".docx":
                    doc = Document(io.BytesIO(data))
                    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
                elif suffix == ".pdf":
                    reader = PdfReader(io.BytesIO(data))
                    text = "\n".join((page.extract_text() or "") for page in reader.pages)
                elif suffix in {".txt", ".html", ".htm", ".xml"}:
                    text = data.decode("utf-8", errors="replace")
                else:
                    continue
                parts.append(f"\nФАЙЛ В АРХИВЕ {name}\n{text}")
            except Exception as error:
                parts.append(f"\nФАЙЛ В АРХИВЕ {name}\n[не удалось извлечь текст: {error}]")
    return "\n".join(parts)


def extract_file(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".docx":
            return extract_docx(path)
        if suffix == ".pdf":
            return extract_pdf(path)
        if suffix in {".xlsx", ".xlsm"}:
            return extract_xlsx(path)
        if suffix == ".zip":
            return extract_zip(path)
        if suffix in {".html", ".htm", ".txt", ".xml"}:
            return path.read_text("utf-8", errors="replace")
        return "[формат требует отдельной конвертации]"
    except Exception as error:
        return f"[не удалось извлечь текст: {error}]"


def page_urls(item: dict[str, object]) -> dict[str, str]:
    registry = str(item["registry"])
    if item["law"] == "223-FZ":
        guid = str(item["guid"])
        return {
            "common": str(item["common_url"]),
            "lots": f"{BASE}/epz/order/notice/notice223/lot-list.html?purchaseNoticeNumber={registry}&noticeGuid={guid}",
            "documents": str(item["documents_url"]),
            "changes": f"{BASE}/epz/order/notice/notice223/changes-and-explanations.html?purchaseNoticeNumber={registry}&noticeGuid={guid}",
            "protocols": f"{BASE}/epz/order/notice/notice223/protocols.html?purchaseNoticeNumber={registry}",
            "contracts": f"{BASE}/epz/order/notice/notice223/contract-info.html?purchaseNoticeNumber={registry}",
            "events": f"{BASE}/epz/order/notice/notice223/event-journal.html?purchaseNoticeNumber={registry}",
        }
    return {
        "common": str(item["common_url"]),
        "documents": str(item["documents_url"]),
        "results": f"{BASE}/epz/order/notice/ea20/view/supplier-results.html?regNumber={registry}",
        "events": f"{BASE}/epz/order/notice/ea20/view/event-journal.html?regNumber={registry}",
        "print": f"{BASE}/epz/order/notice/printForm/view.html?regNumber={registry}",
    }


def main() -> None:
    inventory = {item["registry"]: item for item in json.loads(INVENTORY.read_text("utf-8"))}
    context = ssl.create_default_context(cafile=str(CERT))
    opener = build_opener(HTTPCookieProcessor(CookieJar()), HTTPSHandler(context=context))
    report: list[dict[str, object]] = []

    for registry in SELECTED:
        item = inventory[registry]
        procurement_dir = ROOT / "selected" / registry
        pages_dir = procurement_dir / "pages"
        files_dir = procurement_dir / "files"
        text_dir = procurement_dir / "text"
        pages_dir.mkdir(parents=True, exist_ok=True)
        files_dir.mkdir(parents=True, exist_ok=True)
        text_dir.mkdir(parents=True, exist_ok=True)

        all_page_attachments: list[dict[str, str]] = []
        pages: dict[str, dict[str, str]] = {}
        for page_name, url in page_urls(item).items():
            try:
                value = fetch_html(opener, url, pages_dir / f"{page_name}.html")
                text = html_text(value)
                (text_dir / f"page_{page_name}.txt").write_text(text, "utf-8")
                all_page_attachments.extend(attachment_rows(value))
                pages[page_name] = {"url": url, "status": "ok", "chars": str(len(text))}
            except Exception as error:
                pages[page_name] = {"url": url, "status": f"error: {error}", "chars": "0"}

        attachments = list(item["attachments"]) + all_page_attachments
        unique: list[dict[str, str]] = []
        seen_urls: set[str] = set()
        for attachment in attachments:
            if attachment["url"] in seen_urls:
                continue
            seen_urls.add(attachment["url"])
            unique.append(attachment)
        selected_files = unique if item["law"] == "44-FZ" else select_223_files(unique)

        downloaded: list[dict[str, object]] = []
        for index, attachment in enumerate(selected_files, start=1):
            name = safe_name(attachment["name"])
            if not Path(name).suffix:
                name += ".bin"
            path = files_dir / f"{index:02d}_{name}"
            try:
                if not path.exists() or path.stat().st_size < 100:
                    path.write_bytes(request_bytes(opener, attachment["url"]))
                text = extract_file(path)
                text_path = text_dir / f"file_{index:02d}_{safe_name(Path(name).stem)}.txt"
                text_path.write_text(text, "utf-8")
                downloaded.append({
                    "name": attachment["name"],
                    "url": attachment["url"],
                    "path": str(path),
                    "text_path": str(text_path),
                    "size": path.stat().st_size,
                    "text_chars": len(text),
                })
            except Exception as error:
                downloaded.append({"name": attachment["name"], "url": attachment["url"], "error": str(error)})

        report.append({
            "registry": registry,
            "law": item["law"],
            "pages": pages,
            "downloaded": downloaded,
        })
        print(f"{registry}: {len(pages)} pages, {len(downloaded)} files", flush=True)

    (ROOT / "selected_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), "utf-8")


if __name__ == "__main__":
    main()
